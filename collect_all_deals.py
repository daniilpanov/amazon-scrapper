import os
from json import JSONEncoder, JSONDecoder, JSONDecodeError
from time import sleep
from typing import List

import openpyxl
import requests
from bs4 import BeautifulSoup
from selenium.common import TimeoutException, NoSuchElementException
from selenium.webdriver.common.by import By
from urllib.parse import urlparse, parse_qs

from functions import WebDriver, base_chrome_init

import parser
from helpers import get_all_asins_from_text, log, parse_args


class Level:
    name: str | None
    link: str | None
    items: List['Level']
    asin: str | None
    description: str | None
    image_url: str | None
    ready: bool
    all_items_preloaded: bool
    i: int = 0

    def __init__(self, name, link: str, items=None, asin=None, description=None, img=None, ready=False, all_items_preloaded=False):
        self.name = name
        if link and not link.strip().startswith('https://amazon.com') and not link.strip().startswith('https://www.amazon.com'):
            link = 'https://amazon.com' + link
        self.link = link
        self.asin = asin
        self.description = description
        self.image_url = img
        self.ready = ready
        self.all_items_preloaded = all_items_preloaded
        self.items = []
        if items:
            self.set_items(items)

    def __getitem__(self, item):
        return self.items[item]

    def __setitem__(self, key, value):
        self.items[key] = value

    def __iter__(self):
        self.i = 0
        return self

    def __next__(self):
        if self.i >= len(self.items):
            raise StopIteration
        self.i += 1
        return self.items[self.i - 1]

    def __str__(self):
        if self.asin:
            return (f'Product "{self.name}": asin - {self.asin}, description - {self.description}' 
                    f', image_url - {self.image_url} [{self.link}] {self.ready}')
        return f'Deal "{self.name}" with {len(self.items or [])} items [{self.link}] {self.ready}'

    def __hash__(self):
        return hash(str(self))

    def __len__(self):
        return len(self.items or [])

    def set_items(self, items: List['Level']):
        self.items.extend(items)

    def add_item(self, item: 'Level'):
        self.items.append(item)

    def to_list(self):
        res = []
        for item in [self.name, self.link, self.asin, self.description, self.image_url]:
            if item:
                res.append(item)

        return res

    def to_dict(self):
        items = []
        if self.items:
            for item in self.items:
                items.append(item.to_dict())

        return {'name': self.name, 'link': self.link} | ({
            'asin': self.asin,
            'description': self.description,
            'image_url': self.image_url,
        } if self.asin else {'items': items}) | {'ready': self.ready, 'all_items_preloaded': self.all_items_preloaded}

    @staticmethod
    def from_dict(d):
        level = Level(d.get('name'), d.get('link'), None,
                      d.get('asin'), d.get('description'), d.get('image_url'), d.get('ready', False),
                      d.get('all_items_preloaded', False))
        items = []
        for item in d.get('items', []):
            items.append(Level.from_dict(item))
        level.set_items(items)
        return level

    def add_product_data(self, wd: WebDriver):
        # расширяем данные - добавляем информацию о товаре
        if not self.asin or self.ready:  # проверка: является ли объект товаром и готов ли товар
            return False
        # переходим по ссылке
        print(self.link)
        wd.get(self.link)
        sleep(.1)
        wd.wait_for_loading()
        # парсим tiny набор данных - title, description, image_url
        parsed_data = parser.parse_product(self.asin, wd.get_page_source(), True)
        if not parsed_data:  # если возникла ошибка - возвращаем False
            return False
        self.name, self.description, self.image_url = parsed_data
        return True

    # Фабричный метод
    @staticmethod
    def create_by_a(el, auto_parse=False, wd=None):
        if type(el) not in (tuple, list):
            link = prepare_link(el)
            name = el.text.strip()
        else:
            name, link = el
        asin = get_all_asins_from_text(link)
        level = Level(name, link, asin=asin[0] if len(asin) > 0 else None)
        if auto_parse and wd:
            level.add_product_data(wd)
        return level


class XSheet:
    wb: openpyxl.Workbook
    current_sheet_name: str = None
    current_row: int = 1
    current_subcategory_row: int = 1
    deals_tree: Level
    jsd: JSONDecoder
    jse: JSONEncoder

    def __init__(self, name):
        self.jsd = JSONDecoder()
        self.jse = JSONEncoder()
        self.deals_tree = self.load(name) or Level(name, None)
        path = os.path.join('tmp__', f'{name}.xlsx')
        if os.path.exists(path):
            self.wb = openpyxl.load_workbook(path)
        else:
            self.wb = openpyxl.Workbook()

    def set_link(self, link):
        self.deals_tree.link = link

    @property
    def name(self):
        return self.deals_tree.name

    def save(self):
        if not os.path.isdir('tmp__'):
            os.mkdir('tmp__')
        with open(f'tmp__/{self.name}.json', 'w', encoding='utf-8') as f:
            f.write(self.jse.encode(self.deals_tree.to_dict()))

    def load(self, name):
        if not os.path.exists(f'tmp__/{name}.json'):
            return None
        with open(f'tmp__/{name}.json', encoding='utf-8') as f:
            try:
                d = self.jsd.decode(f.read())
            except JSONDecodeError as e:
                print(e)
                return None
        return Level.from_dict(d)

    def add_sheet(self, name, link):
        self.wb.create_sheet(name)
        self.current_sheet_name = name
        self.current_row = 1
        self.set_cells([link])
        self.current_subcategory_row = 1

    def set_cells(self, cols_vals: list[str | int], marked=False, row=None):
        cols = 'ABCDEFGHJKLMNOPQR'
        if row is None:
            row = self.current_row
        for col_num in range(len(cols_vals)):
            col = cols[col_num]
            self.wb[self.current_sheet_name][col + str(row)] = cols_vals[col_num]
        if marked:
            for i in range(1, len(cols) + 1):
                self.wb[self.current_sheet_name].cell(column=i, row=row).fill = openpyxl.styles.PatternFill(
                    start_color='ffff00',
                    end_color='ffff00',
                    fill_type='solid',
                )
        self.current_row += 1

    def to_xlsx(self):
        self.wb.save(os.path.abspath(os.path.join('tmp__', f'{self.name}.xlsx')))


######################################################################
def search_deals_link(wd: WebDriver):
    wd.get('https://amazon.com')
    try:
        link = wd.find_text('See all deals', timeout=1)
        return link.get_attribute('href')
    except:
        return None


# Функция для получения списка категорий deals
def get_all_deals_categories(wd: WebDriver, deals_tree: Level):
    wd.get(deals_tree.link)
    soup = BeautifulSoup(wd.get_page_source(), features='html.parser')
    deals_els = soup.select('.a-carousel-card[class*="GridPresets-module__gridPresetElement_"]>a')
    links = set(deal.link for deal in deals_tree)
    links_to_skip = set(deal.link for deal in deals_tree if deal.ready)
    for deal_el in deals_els:
        try:
            card_items = deal_el.find_all('span')
            if not card_items or len(card_items) < 2:
                continue
            title = card_items[-1].text.strip()
            if title == 'All Deals':
                continue
            link = prepare_link(deal_el)
            if 'goldbox' not in link or link in links_to_skip:
                continue
        except Exception as e:
            log(e)
            continue
        if link not in links:
            deals_tree.add_item(Level(title, link))
    deals_tree.all_items_preloaded = True
    return deals_tree.items

def get_products_from_deal(wd: WebDriver, deal: Level, xsheet: XSheet):
    log('[3] Goto deal', deal.link)
    # переходим по ссылке deal
    wd.get(deal.link)
    if not deal.all_items_preloaded:
        try:
            # ожидаем загрузки списка asins
            wd.wait_for_loading('.octops-dlp-asin-stream-section, '
                                '#productInfoList,'
                                'span[data-component-type="s-search-results"],'
                                '[class*="_octopus-search-result-card_style_apbSearchResultsContainer__"]', 2)
            log('[3] List of ASINs is loaded')
        except TimeoutException:
            log('ERROR when loading ASINs list:', deal.link)
            return None
        soup = BeautifulSoup(wd.get_page_source(), features='html.parser')
        # получаем все ссылки из списка asins
        links = soup.select(
            '.octops-dlp-asin-stream-section a[href*="B0"],'
            '#productInfoList a[href*="B0"],'
            'span[data-component-type="s-search-results"] a[href*="B0"],'
            '[class*="_octopus-search-result-card_style_apbSearchResultsContainer__"]',
        )
        prepared_links = set(prod.link for prod in deal)
        # перебираем
        for link in links:
            text_link = prepare_link(link)
            if text_link not in prepared_links:
                # добавляем созданный из ссылки уровень
                deal.add_item(Level.create_by_a(link))
        deal.all_items_preloaded = True
        xsheet.save()
    log('[3] Links list:', deal)
    # после получения перебираем уровни и добавляем в каждый информацию о товаре
    for link in deal:
        if link.ready:
            continue
        link.add_product_data(wd)
        link.ready = True
        xsheet.save()
    deal.ready = True
    log('[3] Product data was added')


def get_all_deals_from_category(wd: WebDriver, cat_level: Level, xsheet: XSheet):
    log('[2] Goto category link:', cat_level.link)
    # переход на страницу категории
    wd.get(cat_level.link)
    links = set(deal.link for deal in cat_level)
    i = 0
    soup = BeautifulSoup(wd.get_page_source(), features='html.parser')
    if not cat_level.all_items_preloaded:
        while True:
            log('[2] Loaded! Current page:', i)
            # получаем все карточки deals текущей категории
            deals_els = soup.select('[class*="DealGridItem-module__dealItemDisplayGrid_"]')
            # перебираем
            for deal_el in deals_els:
                # получаем ссылку на текущую deal (последнюю, т.к. в ней содержится название)
                link_el = deal_el.find_all('a')
                if link_el:
                    name = link_el[-1].text.strip()
                    link = prepare_link(link_el[-1])
                else:
                    log('[2] Error when trying to get all "a" elements from deal:', link_el)
                    continue
                log('[2] Deal found:', link)
                if link not in links:
                    log('[2] Add deal:', link)
                    cat_level.add_item(Level.create_by_a((name, link)))  # создаём уровень через ссылку
            # если следующей страницы нет - drop cycle
            try:
                if wd.find_element(by=By.CSS_SELECTOR, value='li.a-last.a-disabled'):
                    log('[2] All pages collected!')
                    break
            except NoSuchElementException:
                pass
            # сохраняемся
            xsheet.save()
            # переходим на следующую страницу
            last_page_url = wd.current_url
            wd.click('li.a-last')
            while True:
                wd.wait_for_loading()
                sleep(3)
                if wd.current_url != last_page_url:
                    break
                wd.wait_for_loading()
                sleep(1)
                wd.click('li.a-last')
            wd.wait_for_loading()
            log('[2] Goto next page -> -> ->')
            if i > 9:
                break
            i += 1
        # все ссылки загружены
        cat_level.all_items_preloaded = True
        xsheet.save()
    log('[2] Foreach levels - add the product data')
    # перебираем полученные данные
    for deal in cat_level:
        if deal.ready:
            continue
        # если deal - товар, то добавляем информацию
        if deal.asin:
            log('[2] Add the additional data to the', deal.name)
            deal.add_product_data(wd)
            deal.ready = True
        else:
            # иначе добавляем элементы - товары
            get_products_from_deal(wd, deal, xsheet)
            log('[2] Deal upgraded:', deal.name)
        xsheet.save()


def collect_all_info(xsheet: XSheet):
    log('[1] Init chrome')
    wd = base_chrome_init(goto='https://amazon.com')
    log('[1] Change loc')
    wd.change_loc()
    wd.wait_for_loading()
    # ищем ссылку на all deals
    if not xsheet.deals_tree.link:
        deals_link = prepare_link(search_deals_link(wd))
        if not deals_link:
            return False
        xsheet.set_link(deals_link)
        log('[1] All Deals link found:', deals_link)
    # ищем все категории deals, если не собраны
    if xsheet.deals_tree.all_items_preloaded:
        deals_categories = xsheet.deals_tree.items
    else:
        deals_categories = get_all_deals_categories(wd, xsheet.deals_tree)
        xsheet.save()
    log('[1] Deals categories found:', xsheet.deals_tree)
    # перебираем
    for category in deals_categories:
        if category.ready:
            continue
        # добавляем элементы - deals (see this function)
        log('[1] Get all deals from category:', category.name)
        get_all_deals_from_category(wd, category, xsheet)
        category.ready = True
        xsheet.save()
    xsheet.deals_tree.ready = True
    xsheet.save()


def convert_from_excel_format(cell: str):
    conv_part = ''
    row = ''
    for c in cell:
        if c.isdigit():
            row += c
        else:
            conv_part += c
    col = 0
    for i in range(len(conv_part)):
        col += (ord(conv_part[len(conv_part) - 1 - i]) - ord('A') + 1) * 26 ** i
    return int(row), col


def add_marked_cell(wb, sheet, cell: str, value):
    row, col = convert_from_excel_format(cell)
    sheet[cell] = value
    wb.active.cell(column=col, row=row).fill = openpyxl.styles.PatternFill(
        start_color='ffff00',
        end_color='ffff00',
        fill_type='solid',
    )


def write_info(xsheet: XSheet):
    if not os.path.exists('tmp__'):
        os.mkdir('tmp__')
    log('[1] Write lists')
    for list in xsheet.deals_tree:
        xsheet.add_sheet(list.name, list.link)
        log('[1] List writen:', list.name, '\n', list.link)
        for item in list:
            xsheet.set_cells(item.to_list(), True)
            if not item.asin:
                i = 1
                for asin in item:
                    xsheet.set_cells([i] + asin.to_list())
                    i += 1
            log('[2] Deal writen:', item)
    xsheet.to_xlsx()


def prepare_link(link):
    if type(link) is not str:
        link = link['href']
    parsed_link = urlparse(link)
    query = parse_qs(parsed_link.query or '')
    new_query = ('?deals-widget=' + query.get('deals-widget', [''])[0]) if 'deals-widget' in query else ''
    return link.replace('?' + parsed_link.query, new_query)


def start(name=None, tg_note_users_ids: list[str] | None = None):
    if not tg_note_users_ids:
        tg_note_users_ids = []
    xsh = XSheet(name or hash(tg_note_users_ids))
    if not xsh.deals_tree.ready:
        log('[0] Start. Collect all info')
        collect_all_info(xsh)
    log('[0] Write all info')
    write_info(xsh)
    log('[0] Open file to send it')
    with open(os.path.join('tmp__', name + '.xlsx'), 'rb') as f:
        requests.post(
            'localhost:8080',
            {'msg': 'Data collected! Your XLSX file with the Deals:', 'uid': ','.join(tg_note_users_ids)},
            files=[(name + '.xlsx', f)],
        )
        log('[0] Sent!')
    log('[0] Program finished')


if __name__ == '__main__':
    import sys
    params_dict = parse_args(sys.argv)
    try:
        start(params_dict.get('sheet_name', 'TMP'), params_dict.get('user', '1456674317,1428909514').split(','))
        requests.post('http://localhost:8080/send_msg', {
            'msg': f'Deals collected: {params_dict.get("sheet_name")}',
            'uid': params_dict['user'],
        })
    except Exception as e:
        requests.post('http://localhost:8080/send_msg', {
            'msg': f'Error on collecting deals: {params_dict.get("sheet_name")}\n' + str(e) + '\n',
            'uid': params_dict['user'],
        })
    finally:
        requests.post('http://localhost:8080/end_task', {'_id': params_dict['_id']})
