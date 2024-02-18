import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from json import JSONDecoder, JSONEncoder, JSONDecodeError
from time import sleep
from typing import List

import openpyxl
from openpyxl.utils.cell import get_column_letter
import requests
from bs4 import BeautifulSoup

from functions import WebDriver, base_chrome_init
from helpers import log, parse_args, send_bot_msg, path


domain = 'amazon.com'


class Level:
    name: str | None
    link: str | None
    items: List['Level']
    is_last_group: bool
    ready: bool
    all_items_preloaded: bool
    i: int = 0

    def __init__(
            self, name, link: str | None = None, items=None,
            is_last_group=None, ready: bool = False, all_items_preloaded: bool = False,
    ):
        self.name = name
        if (link and not link.strip().startswith(f'https://{domain}')
                and not link.strip().startswith(f'https://www.{domain}')):
            link = f'https://{domain}' + link
        self.link = link
        self.is_last_group = is_last_group
        self.ready = ready
        self.all_items_preloaded = all_items_preloaded
        self.items = []
        if items:
            self.set_items(items)

    def __getitem__(self, item):
        return self.items[item]

    def __setitem__(self, key, value):
        self.items[key] = value

    def __iter__(self):
        self.i = 0
        return self

    def __next__(self):
        if self.i >= len(self.items):
            raise StopIteration
        self.i += 1
        return self.items[self.i - 1]

    def __str__(self):
        return f'Department "{self.name}" with {len(self.items or [])} items [{self.link}] {self.ready}'

    def __hash__(self):
        return hash(str(self))

    def __len__(self):
        return len(self.items or [])

    def set_items(self, items: List['Level']):
        self.items.extend(items)

    def check_exists(self, name):
        for item in self.items:
            if name == item.name:
                return True
        return False

    def add_item(self, item: 'Level', unique=False):
        if not unique or not self.check_exists(item.name):
            self.items.append(item)

    def to_list(self):
        res = []
        for item in [self.name, self.link]:
            if item:
                res.append(item)

        return res

    def to_dict(self):
        items = []
        if self.items:
            for item in self.items:
                items.append(item.to_dict())

        return {
            'name': self.name, 'link': self.link,
            'is_last_group': self.is_last_group,
            'items': items, 'ready': self.ready,
            'all_items_preloaded': self.all_items_preloaded,
        }

    @staticmethod
    def from_dict(d):
        level = Level(d.get('name'), d.get('link'), None, d.get('is_last_group', False),
                      d.get('ready', False), d.get('all_items_preloaded', False))
        items = []
        for item in d.get('items', []):
            items.append(Level.from_dict(item))
        level.set_items(items)
        return level


class XSheet:
    wb: openpyxl.Workbook
    current_row: int = 1
    departments_tree: Level
    jsd: JSONDecoder
    jse: JSONEncoder
    browsers: list[list[bool, WebDriver]]
    workers: int

    def __init__(self, name, workers=1):
        self.jsd = JSONDecoder()
        self.jse = JSONEncoder()
        self.departments_tree = self.load(name) or Level(name)
        p = path('tmp__', f'{name}.xlsx', filecontent=False)
        if os.path.exists(p):
            self.wb = openpyxl.load_workbook(p)
        else:
            self.wb = openpyxl.Workbook()
        self.wb.active.title = name
        self.browsers = []
        self.workers = workers

    def add_browser(self, wd: WebDriver):
        self.browsers.append([True, wd])
        return len(self.browsers) - 1

    def get_browser(self):
        for i in range(len(self.browsers)):
            # Если браузер занят
            if not self.browsers[i][0]:
                continue
            self.browsers[i][0] = False
            return self.browsers[i][1], i
        return None, -1

    def return_browser(self, i):
        self.browsers[i][0] = True

    def set_link(self, link):
        self.departments_tree.link = link
        return link

    @property
    def name(self):
        return self.departments_tree.name

    def save(self):
        with open(path('tmp__', f'{self.name}.json', filecontent=False), 'w', encoding='utf-8') as f:
            f.write(self.jse.encode(self.departments_tree.to_dict()))

    def load(self, name):
        if not os.path.exists(f'tmp__/{name}.json'):
            return None
        with open(f'tmp__/{name}.json', encoding='utf-8') as f:
            try:
                d = self.jsd.decode(f.read())
            except JSONDecodeError as e:
                print(e)
                return None
        return Level.from_dict(d)

    def set_cells(self, cols_vals: list[str | int]):
        for col_num in range(len(cols_vals)):
            self.wb.active.cell(self.current_row, col_num + 1).value = cols_vals[col_num]
            if self.current_row > 1:
                w = max(
                    self.wb.active.column_dimensions[get_column_letter(col_num + 1)].width,
                    len(cols_vals[col_num]) * 11**(11*0.009),
                )
            else:
                w = len(cols_vals[col_num]) * 11**(11*0.009)
            self.wb.active.column_dimensions[get_column_letter(col_num + 1)].width = w
        self.current_row += 1

    def to_xlsx(self):
        self.wb.save(os.path.abspath(path('tmp__', f'{self.name}.xlsx', filecontent=False)))


def recursive_tree(tree: Level, xsh: XSheet, name_only=None):
    if tree.ready:
        return
    wd, wd_id = xsh.get_browser()
    while not wd:
        sleep(1)
        wd, wd_id = xsh.get_browser()
    # if links are not collected
    if not tree.all_items_preloaded:
        wd.get(tree.link)
        soup = BeautifulSoup(wd.get_page_source(), features='html.parser')
        group = soup.find(role='tree').find(role='group')
        tree.is_last_group = bool(group.select('[role="treeitem"] > span') if group else False)
        if not tree.is_last_group:
            links_a = group.find_all('a')
            for a_tag in links_a or []:
                name = a_tag.text
                if name:
                    name = name.strip()
                log('[2] Departments found:', name)
                if not name_only or name_only == name:
                    tree.add_item(Level(name, a_tag['href']), True)
        tree.all_items_preloaded = True
        xsh.save()
    if wd_id is not None:
        xsh.return_browser(wd_id)
    # if links are already collected, or we have a links group
    if tree.is_last_group is False:
        # RECURSIVE_CALL: load tree
        # full loading
        return (link for link in tree if not name_only or name_only == link.name)


def recursive_call(pool: ThreadPoolExecutor, trees, xsh):
    futures = []
    for tree in trees:
        futures.append(pool.submit(recursive_tree, tree, xsh))
    for future in as_completed(futures):
        trees = future.result()
        if trees:
            recursive_call(pool, trees, xsh)


def collect_all_info(xsheet: XSheet, dep_name=None):
    log('[1] Get chrome')
    xsheet.set_link(f'https://{domain}/gp/bestsellers')
    with ThreadPoolExecutor(xsheet.workers) as pool:
        recursive_tree(xsheet.departments_tree, xsheet, dep_name)
        recursive_call(pool, xsheet.departments_tree.items, xsheet)
    xsheet.save()
    log('[1] Tree collected:', xsheet.departments_tree)
    xsheet.departments_tree.ready = True
    xsheet.save()


def write_info(xsh: XSheet, tree: Level, cat_names: list[str] | None = None):
    if not cat_names:
        cat_names = []
    if tree.is_last_group:
        xsh.set_cells(cat_names + [tree.name])
    for item in tree:
        log(item)
        write_info(xsh, item, cat_names + [tree.name])
    return True


def add_browser():
    wd = base_chrome_init(goto=f'https://{domain}')
    wd.change_loc(domain=domain)
    return wd


def start(sheet_name=None, dep_name=None, tg_note_user_id: int | str | None = True, workers=1):
    if not sheet_name:
        sheet_name = str(hash(tg_note_user_id))
    xsh = XSheet(sheet_name, workers)

    with ThreadPoolExecutor(os.cpu_count()) as pool:
        for _ in range(workers):
            fut = pool.submit(add_browser)
            fut.add_done_callback(lambda res: xsh.add_browser(res.result()) if res.done() else None)

    if not xsh.departments_tree.ready:
        log('[0] Start. Collect all info')
        collect_all_info(xsh, dep_name)
    if tg_note_user_id:
        log('[0] Write all info')
        for item in xsh.departments_tree:
            write_info(xsh, item)
        xsh.to_xlsx()
        log('[0] Open file to send it')
        with open(path('tmp__', sheet_name + '.xlsx'), 'rb') as f:
            send_bot_msg(
                tg_note_user_id,
                'Data collected! Your XLSX file with the Departments:',
                [(sheet_name + '.xlsx', f)],
            )
            log('[0] Sent!')
    log('[0] Program finished')


if __name__ == '__main__':
    import sys
    params_dict = parse_args(sys.argv)
    domain = params_dict.get('domain', 'amazon.com')
    try:
        params_dict.setdefault('dep_name', None)
        # params_dict.setdefault('user', '1428909514')
        params_dict.setdefault('user', '1456674317')
        start(
            params_dict['dep_name'] or 'all departments',
            params_dict['dep_name'], params_dict['user'], 2,
        )
        send_bot_msg(params_dict['user'], f'Departments collected: {params_dict["dep_name"]}')
    except Exception as e:
        send_bot_msg(
            params_dict['user'],
            f'Error on collecting departments: {params_dict["dep_name"]}\n' + str(e) + '\n',
        )
    finally:
        if '_id' in params_dict:
            requests.post('http://localhost:8080/end_task', {'_id': params_dict['_id']})
