import os

import openpyxl

wb = openpyxl.Workbook()

sheet = wb['Sheet']
sheet.title = 'Test List'
sheet['A1'] = 'https://amazon.com'

sheet[f'A2'] = 'NAME'
wb.active.cell(column=1, row=2).fill = openpyxl.styles.PatternFill(
    start_color='ffff00',
    end_color='ffff00',
    fill_type='solid',
)
sheet[f'B2'] = 'LINK'
wb.active.cell(column=2, row=2).fill = openpyxl.styles.PatternFill(
    start_color='ffff00',
    end_color='ffff00',
    fill_type='solid',
)
is_asin = False
if is_asin:
    sheet[f'C2'] = 'ASIN'
    wb.active.cell(column=3, row=2).fill = openpyxl.styles.PatternFill(
        start_color='ffff00',
        end_color='ffff00',
        fill_type='solid',
    )
    sheet[f'D2'] = 'TITLE'
    wb.active.cell(column=4, row=2).fill = openpyxl.styles.PatternFill(
        start_color='ffff00',
        end_color='ffff00',
        fill_type='solid',
    )
    sheet[f'E2'] = 'DESCRIPTION'
    wb.active.cell(column=5, row=2).fill = openpyxl.styles.PatternFill(
        start_color='ffff00',
        end_color='ffff00',
        fill_type='solid',
    )
    sheet[f'F2'] = 'IMAGE_URL'
    wb.active.cell(column=6, row=2).fill = openpyxl.styles.PatternFill(
        start_color='ffff00',
        end_color='ffff00',
        fill_type='solid',
    )
else:
    sheet[f'A3'] = 1
    sheet[f'B3'] = 'LINK'
    sheet[f'C3'] = 'ASIN'
    sheet[f'D3'] = 'TITLE'
    sheet[f'E3'] = 'DESCRIPTION'
    sheet[f'F3'] = 'IMAGE_URL'
    sheet[f'A4'] = 2
    sheet[f'B4'] = 'LINK'
    sheet[f'C4'] = 'ASIN'
    sheet[f'D4'] = 'TITLE'
    sheet[f'E4'] = 'DESCRIPTION'
    sheet[f'F4'] = 'IMAGE_URL'

wb.save(os.path.join('tmp', 'test.xlsx'))